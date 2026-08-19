"""biz/writer/render.py — docx 渲染器（M3）。

实现平台 ArtifactRenderer 协议：把合并稿 md 渲染为带模板样式的 Word 文档。

- mini md→docx：标题(#/##/###)、表格(| |)、列表(-)、粗体(** **)、段落
- 模板样式：标题字体/字号/颜色、页眉，由 TEMPLATE_STYLE 定义（业务可扩展）
- 输出：{session_dir}/output/方案_v{version}.docx
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# 模板样式（M3 最小集：标题层级 + 页眉；图表/页脚样式留后续扩展）
TEMPLATE_STYLE: dict = {
    "h1": {"size": 18, "bold": True, "color": (0x1F, 0x3B, 0x73)},   # 深蓝
    "h2": {"size": 15, "bold": True, "color": (0x2E, 0x59, 0xA8)},   # 中蓝
    "h3": {"size": 13, "bold": True, "color": (0x44, 0x44, 0x44)},   # 深灰
    "body": {"size": 11, "color": (0x33, 0x33, 0x33)},
    "header_text": "施工方案",                                        # 页眉
}


class DocxRenderer:
    """docx 渲染器：实现 ArtifactRenderer 协议。"""

    name = "docx"

    def __init__(self, style: dict | None = None):
        self.style = style or TEMPLATE_STYLE

    def render(self, session, *, merged: str, outline: str,
               version: int, **kw) -> str:
        doc = Document()

        # 页眉
        header = doc.sections[0].header
        header.paragraphs[0].text = self.style.get("header_text", "")

        # 大纲在前（标题页）, 合并稿在后
        self._add_md(doc, outline)
        doc.add_page_break()
        self._add_md(doc, merged)

        filename = f"方案_v{version}.docx"
        path = os.path.join(session.dir, "output", filename)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        doc.save(path)
        return filename

    # ── mini md → docx ──────────────────────────────────

    def _add_md(self, doc: Document, md: str) -> None:
        lines = md.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip():
                i += 1
                continue

            # 表格（连续 | 行合并为一张表）
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(lines[i].strip())
                    i += 1
                self._add_table(doc, rows)
                continue

            # 标题
            m = re.match(r"^(#{1,6})\s+(.*)$", line)
            if m:
                self._add_heading(doc, m.group(2).strip(), len(m.group(1)))
                i += 1
                continue

            # 列表
            if line.startswith("- ") or line.startswith("* "):
                p = doc.add_paragraph(style="List Bullet")
                self._add_inline(p, line[2:].strip())
                i += 1
                continue

            # 普通段落
            p = doc.add_paragraph()
            self._add_inline(p, line)
            i += 1

    def _add_heading(self, doc: Document, text: str, level: int) -> None:
        level = min(level, 3)
        h = doc.add_heading(level=level)
        run = h.add_run(self._plain(text))
        s = self.style.get(f"h{level}", self.style["body"])
        run.font.size = Pt(s.get("size", 12))
        run.font.bold = s.get("bold", False)
        run.font.color.rgb = RGBColor(*s.get("color", (0x33, 0x33, 0x33)))

    def _add_table(self, doc: Document, rows: list[str]) -> None:
        # 跳过 md 表格分隔行（|---|------|）
        cells = [
            [c.strip() for c in r.strip("|").split("|")]
            for r in rows
            if not all(re.fullmatch(r"-{2,}", c.strip()) for c in r.strip("|").split("|"))
        ]
        if not cells:
            return
        ncols = max(len(r) for r in cells)
        table = doc.add_table(rows=len(cells), cols=ncols)
        for ri, row in enumerate(cells):
            for ci in range(ncols):
                cell = table.cell(ri, ci)
                cell.text = row[ci] if ci < len(row) else ""
                if ri == 0 and cell.paragraphs[0].runs:  # 表头加粗
                    cell.paragraphs[0].runs[0].bold = True

    def _add_inline(self, paragraph, text: str) -> None:
        """行内解析: **粗体** 分段添加 run。"""
        for part in re.split(r"(\*\*.+?\*\*)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            else:
                paragraph.add_run(part)

    @staticmethod
    def _plain(text: str) -> str:
        return text.replace("**", "")
