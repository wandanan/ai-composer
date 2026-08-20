"""platform/extract — 文档提取插件（provides=["extract"]）。

格式: .docx/.doc（Word→PDF→pymupdf; .docx 回退 python-docx）/ .pdf（pymupdf）/ 文本兜底。
引擎: local（pymupdf/docx）/ mineru（可选 API）。换引擎只改本插件, 消费方零改动。
"""
from __future__ import annotations

import io
import logging
import os
from typing import Protocol, runtime_checkable

from aic.kernel import Context, Plugin

from aic.extensions.platform.extract.convert_to_pdf import convert_to_pdf

logger = logging.getLogger(__name__)

ALLOWED_TYPES = {".docx": "docx", ".doc": "docx", ".pdf": "pdf",
                 ".txt": "txt", ".md": "md"}


def extract_document(content: bytes, filename: str, use_ocr: bool = False,
                     progress_callback=None, engine: str = "local") -> str:
    """从文件字节提取文本（Word 自动转 PDF; 引擎 local/mineru）。

    大声失败原则: 不支持的类型/缺依赖/转换失败 → 抛异常（调用方自行 catch）;
    仅"文档本身为空"返回 ""。不要把"提取失败"和"内容为空"混为一谈。
    """
    if progress_callback:
        try:
            progress_callback({"stage": "preprocess", "message": "开始提取"})
        except Exception:
            pass
    ext = os.path.splitext(filename)[1].lower()

    if ext in (".docx", ".doc"):
        # Word → PDF（LibreOffice）→ 提取（对齐原项目 convert_to_pdf 管线）
        try:
            pdf = convert_to_pdf(content, filename)
        except Exception as e:
            logger.warning(f"[extract] Word→PDF 失败 {filename}: {e}, 回退 python-docx")
            if ext == ".docx":
                return _extract_docx(content)
            raise RuntimeError(f"[extract] .doc 转 PDF 失败: {e}") from e
        if engine == "mineru":
            return _extract_mineru(pdf, filename, progress_callback)
        return _extract_pdf(pdf)
    if ext == ".pdf":
        if engine == "mineru":
            return _extract_mineru(content, filename, progress_callback)
        return _extract_pdf(content)
    if ext in (".txt", ".md"):
        return content.decode("utf-8", errors="replace")
    raise ValueError(f"[extract] 不支持的文件类型: {ext} ({filename})")


def _extract_mineru(content: bytes, filename: str, progress_callback=None) -> str:
    """MinerU 引擎（批量提交+轮询+MD5 缓存, 对齐原项目 mineru_extract）。"""
    import os

    from aic.extensions.platform.extract.mineru_client import MinerUClient

    client = MinerUClient(
        base_url=os.environ.get("MINERU_API_URL", "http://127.0.0.1:8000"),
        api_key=os.environ.get("MINERU_API_KEY", ""),
        cache_enabled=os.environ.get("MINERU_CACHE_ENABLED", "1") == "1")
    return client.extract(content, filename, progress_callback)


def _extract_docx(content: bytes) -> str:
    """python-docx 提取段落 + 表格文本（缺依赖/解析失败 → RuntimeError）。"""
    try:
        import docx
    except ImportError:
        raise RuntimeError("[extract] python-docx 未安装, 无法提取 docx") from None
    try:
        doc = docx.Document(io.BytesIO(content))
        parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        logger.info(f"[extract] docx 提取: {len(text)} 字符")
        return text
    except Exception as e:
        raise RuntimeError(f"[extract] docx 提取失败: {e}") from e


def _extract_pdf(content: bytes) -> str:
    """pymupdf 逐页提取文本（缺依赖/解析失败 → RuntimeError）。"""
    try:
        import fitz
    except ImportError:
        raise RuntimeError("[extract] pymupdf 未安装, 无法提取 pdf") from None
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        parts = [page.get_text() for page in doc]
        text = "\n".join(parts)
        logger.info(f"[extract] pdf 提取: {len(text)} 字符 / {len(doc)} 页")
        return text
    except Exception as e:
        raise RuntimeError(f"[extract] pdf 提取失败: {e}") from e


@runtime_checkable
class FileExtractor(Protocol):
    """文档提取协议（消费方依赖此形状, 不 import 实现）。"""

    def extract(self, content: bytes, filename: str, use_ocr: bool = False,
                progress_callback=None, engine: str = "local") -> str: ...


class LocalExtractor:
    """本地提取引擎（pymupdf/docx + Word→PDF 回退）。"""

    def extract(self, content, filename, use_ocr=False, progress_callback=None,
                engine="local") -> str:
        return extract_document(content, filename, use_ocr, progress_callback, engine)


class ExtractPlugin(Plugin):
    """文档提取插件：提供 ctx.extract。构造注入实现（默认 LocalExtractor）。"""

    PUBLIC = True   # 公共插件: 有独立生命周期, 不随任何应用卸载删除
    provides = ["extract"]

    def __init__(self, impl=None):
        self._impl = impl

    def apply(self, ctx: Context):
        ctx.register("extract", self._impl or LocalExtractor())
